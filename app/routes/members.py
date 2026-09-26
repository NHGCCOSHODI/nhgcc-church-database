import io
from datetime import datetime, date
from flask import Blueprint, render_template, request, redirect, url_for, flash, jsonify, send_file
from sqlalchemy import or_
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from app.database import db
from app.models import Member, Attendance, FollowUpLog
from app.services.notification_logic import get_recent_sundays

members_bp = Blueprint('members', __name__, url_prefix='/members')

@members_bp.route('/')
def index():
    search_q = request.args.get('q', '').strip()
    dept_filter = request.args.get('department', '').strip()
    status_filter = request.args.get('status', '').strip()
    gender_filter = request.args.get('gender', '').strip()
    month_filter = request.args.get('month', '').strip()

    query = Member.query.filter_by(is_active=True)

    if search_q:
        pat = f"%{search_q}%"
        query = query.filter(
            or_(
                Member.full_name.ilike(pat),
                Member.phone.ilike(pat),
                Member.email.ilike(pat),
                Member.address.ilike(pat),
                Member.member_code.ilike(pat)
            )
        )

    if dept_filter:
        query = query.filter_by(department=dept_filter)
    if status_filter:
        query = query.filter_by(status=status_filter)
    if gender_filter:
        query = query.filter_by(gender=gender_filter)
    if month_filter and month_filter.isdigit():
        query = query.filter_by(birth_month=int(month_filter))

    members = query.order_by(Member.id).all()

    dept_rows = db.session.query(Member.department).distinct().filter(Member.department != None).all()
    departments = [d[0] for d in dept_rows if d[0]]

    return render_template(
        'members/index.html',
        members=members,
        total_count=len(members),
        search_q=search_q,
        dept_filter=dept_filter,
        status_filter=status_filter,
        gender_filter=gender_filter,
        month_filter=month_filter,
        departments=departments
    )

@members_bp.route('/<int:member_id>')
def profile(member_id):
    member = Member.query.get_or_404(member_id)

    recent_attendances = Attendance.query.filter_by(member_id=member.id).order_by(
        Attendance.service_date.desc()
    ).limit(12).all()

    total_attended = Attendance.query.filter_by(member_id=member.id, status='Present').count()

    last_4_sundays = get_recent_sundays(count=4)
    attended_in_last_4 = Attendance.query.filter(
        Attendance.member_id == member.id,
        Attendance.service_date.in_(last_4_sundays),
        Attendance.status == 'Present'
    ).count()

    retention_rate = int((attended_in_last_4 / 4) * 100) if last_4_sundays else 0

    follow_ups = FollowUpLog.query.filter_by(member_id=member.id).order_by(
        FollowUpLog.contact_date.desc()
    ).all()

    return render_template(
        'members/profile.html',
        member=member,
        recent_attendances=recent_attendances,
        total_attended=total_attended,
        retention_rate=retention_rate,
        follow_ups=follow_ups
    )

@members_bp.route('/new', methods=['GET', 'POST'])
def create():
    if request.method == 'POST':
        full_name = request.form.get('full_name', '').strip()
        title = request.form.get('title', 'Member')
        phone = request.form.get('phone', '').strip()
        email = request.form.get('email', '').strip()
        gender = request.form.get('gender', 'Not Specified')
        address = request.form.get('address', '').strip()
        department = request.form.get('department', 'Congregation')
        status = request.form.get('status', 'Active')
        dob_str = request.form.get('date_of_birth', '').strip()
        notes = request.form.get('notes', '').strip()

        if not full_name:
            flash('Member name is required.', 'danger')
            return render_template('members/form.html', member=None)

        last_m = Member.query.order_by(Member.id.desc()).first()
        next_id = (last_m.id + 1) if last_m else 1
        member_code = f"NHGCC-{next_id:04d}"

        dob = None
        bday = None
        bmonth = None
        if dob_str:
            try:
                dob = datetime.strptime(dob_str, '%Y-%m-%d').date()
                bday = dob.day
                bmonth = dob.month
            except ValueError:
                pass

        parts = full_name.split(maxsplit=1)
        first_name = parts[0]
        last_name = parts[1] if len(parts) > 1 else ''

        new_m = Member(
            member_code=member_code,
            full_name=full_name,
            first_name=first_name,
            last_name=last_name,
            title=title,
            phone=phone,
            email=email,
            gender=gender,
            address=address,
            department=department,
            status=status,
            date_of_birth=dob,
            birth_day=bday,
            birth_month=bmonth,
            notes=notes,
            date_joined=date.today()
        )
        db.session.add(new_m)
        db.session.commit()

        flash(f"Member '{new_m.full_name}' was successfully registered!", 'success')
        return redirect(url_for('members.profile', member_id=new_m.id))

    return render_template('members/form.html', member=None)

@members_bp.route('/<int:member_id>/edit', methods=['GET', 'POST'])
def edit(member_id):
    member = Member.query.get_or_404(member_id)

    if request.method == 'POST':
        member.full_name = request.form.get('full_name', '').strip()
        member.title = request.form.get('title', 'Member')
        member.phone = request.form.get('phone', '').strip()
        member.email = request.form.get('email', '').strip()
        member.gender = request.form.get('gender', 'Not Specified')
        member.address = request.form.get('address', '').strip()
        member.department = request.form.get('department', 'Congregation')
        member.status = request.form.get('status', 'Active')
        dob_str = request.form.get('date_of_birth', '').strip()
        member.notes = request.form.get('notes', '').strip()

        if dob_str:
            try:
                dob = datetime.strptime(dob_str, '%Y-%m-%d').date()
                member.date_of_birth = dob
                member.birth_day = dob.day
                member.birth_month = dob.month
            except ValueError:
                pass
        else:
            member.date_of_birth = None
            member.birth_day = None
            member.birth_month = None

        db.session.commit()
        flash(f"Details for '{member.full_name}' updated successfully!", 'success')
        return redirect(url_for('members.profile', member_id=member.id))

    return render_template('members/form.html', member=member)

@members_bp.route('/<int:member_id>/delete', methods=['POST'])
def delete(member_id):
    member = Member.query.get_or_404(member_id)
    name = member.full_name
    db.session.delete(member)
    db.session.commit()
    flash(f"Member '{name}' has been removed.", 'info')
    return redirect(url_for('members.index'))

@members_bp.route('/export')
def export_excel():
    members = Member.query.filter_by(is_active=True).order_by(Member.id).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'NHGCC Members'

    title_font = Font(name='Segoe UI', size=15, bold=True, color='1E3E62')
    ws.append(['NHGCC OSHODI - OFFICIAL MEMBERS DIRECTORY'])
    ws.append([f"Total Registered Members: {len(members)} | Generated on {date.today().strftime('%d %B %Y')}"])
    ws.append([])

    ws['A1'].font = title_font
    ws['A2'].font = Font(name='Segoe UI', size=11, italic=True)

    headers = [
        'S/N', 'Member Code', 'Full Name', 'Title', 'Gender',
        'Phone Number', 'Email Address', 'Department', 'Status',
        'Birthday', 'House Address'
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color='0B192C', end_color='0B192C', fill_type='solid')
    header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=4, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font

    for idx, m in enumerate(members, 1):
        ws.append([
            idx,
            m.member_code or '',
            m.full_name,
            m.title or '',
            m.gender or '',
            m.phone or '',
            m.email or '',
            m.department or '',
            m.status or '',
            m.formatted_birthday,
            m.address or ''
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"NHGCC_Oshodi_Members_Directory_{date.today().strftime('%Y%m%d')}.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
