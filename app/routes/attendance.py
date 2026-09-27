import io
from datetime import date, datetime
from flask import Blueprint, render_template, request, jsonify, send_file, flash, redirect, url_for
from sqlalchemy import or_, func
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from app.database import db
from app.models import Member, Attendance
from app.services.notification_logic import get_latest_sunday, get_upcoming_sunday

attendance_bp = Blueprint('attendance', __name__, url_prefix='/attendance')


@attendance_bp.route('/checkin')
def checkin():
    today = date.today()
    if today.weekday() == 6:
        default_sunday = get_latest_sunday(today)
    else:
        default_sunday = get_upcoming_sunday(today)

    date_str = request.args.get('date', default_sunday.strftime('%Y-%m-%d'))
    service_type = request.args.get('service_type', 'Sunday Service')

    try:
        service_date = datetime.strptime(date_str, '%Y-%m-%d').date()
    except ValueError:
        service_date = default_sunday

    members = Member.query.filter_by(is_active=True).order_by(Member.full_name).all()
    attendances = Attendance.query.filter_by(service_date=service_date, service_type=service_type).all()

    attended_member_ids = {a.member_id: a for a in attendances if a.status == 'Present'}
    total_active = len(members)
    total_present = len(attended_member_ids)
    first_timers_present = sum(1 for a in attended_member_ids.values() if a.member and a.member.status == 'First-Timer')

    return render_template(
        'attendance/checkin.html',
        service_date=service_date,
        service_date_str=service_date.strftime('%Y-%m-%d'),
        service_type=service_type,
        members=members,
        attended_member_ids=attended_member_ids,
        total_active=total_active,
        total_present=total_present,
        first_timers_present=first_timers_present
    )


@attendance_bp.route('/api/search')
def api_search():
    """Live search API for entrance check-in desk."""
    query = request.args.get('q', '').strip()
    date_str = request.args.get('date')
    service_type = request.args.get('service_type', 'Sunday Service')

    try:
        service_date = datetime.strptime(date_str, '%Y-%m-%d').date()
    except Exception:
        service_date = get_latest_sunday()

    members_query = Member.query.filter(Member.is_active == True)
    if query:
        pattern = f"%{query}%"
        members_query = members_query.filter(
            or_(
                Member.full_name.ilike(pattern),
                Member.phone.ilike(pattern),
                Member.member_code.ilike(pattern),
                Member.department.ilike(pattern)
            )
        )

    members = members_query.order_by(Member.full_name).limit(50).all()
    member_ids = [m.id for m in members]

    attendances = Attendance.query.filter(
        Attendance.member_id.in_(member_ids),
        Attendance.service_date == service_date,
        Attendance.service_type == service_type
    ).all()

    present_dict = {
        a.member_id: a.check_in_time.strftime('%I:%M %p')
        for a in attendances if a.status == 'Present' and a.check_in_time
    }

    results = []
    for m in members:
        is_present = m.id in present_dict
        results.append({
            'id': m.id,
            'member_code': m.member_code or '',
            'full_name': m.display_title_name,
            'phone': m.phone or '',
            'email': m.email or '',
            'department': m.department or 'Congregation',
            'status': m.status or 'Active',
            'is_present': is_present,
            'check_in_time': present_dict.get(m.id, '')
        })

    return jsonify({
        'success': True,
        'count': len(results),
        'members': results
    })


@attendance_bp.route('/api/toggle', methods=['POST'])
def api_toggle():
    """Toggles a member's check-in status for a given service date."""
    data = request.get_json() or {}
    member_id = data.get('member_id')
    date_str = data.get('service_date')
    service_type = data.get('service_type', 'Sunday Service')

    if not member_id or not date_str:
        return jsonify({'success': False, 'message': 'Missing member ID or date'}), 400

    try:
        service_date = datetime.strptime(date_str, '%Y-%m-%d').date()
    except ValueError:
        return jsonify({'success': False, 'message': 'Invalid date format'}), 400

    member = Member.query.get(member_id)
    if not member:
        return jsonify({'success': False, 'message': 'Member not found'}), 404

    record = Attendance.query.filter_by(
        member_id=member.id,
        service_date=service_date,
        service_type=service_type
    ).first()

    if record:
        db.session.delete(record)
        db.session.commit()
        is_present = False
        check_time = ''
        action_msg = f"Marked {member.full_name} as ABSENT"
    else:
        record = Attendance(
            member_id=member.id,
            service_date=service_date,
            service_type=service_type,
            status='Present',
            check_in_time=datetime.now(),
            marked_by='Entrance Desk'
        )
        db.session.add(record)
        db.session.commit()
        is_present = True
        check_time = record.check_in_time.strftime('%I:%M %p')
        action_msg = f"Checked in {member.full_name} successfully!"

    total_present = Attendance.query.filter_by(
        service_date=service_date,
        service_type=service_type,
        status='Present'
    ).count()

    return jsonify({
        'success': True,
        'is_present': is_present,
        'check_in_time': check_time,
        'message': action_msg,
        'total_present': total_present
    })


@attendance_bp.route('/quick-add', methods=['POST'])
def quick_add():
    """
    Instantly registers a newcomer/unregistered attendee from the check-in modal
    and automatically marks them present for today's service!
    """
    full_name = request.form.get('full_name', '').strip()
    phone = request.form.get('phone', '').strip()
    gender = request.form.get('gender', 'Not Specified')
    address = request.form.get('address', '').strip()
    status = request.form.get('status', 'First-Timer')
    department = request.form.get('department', 'Congregation')
    email = request.form.get('email', '').strip()
    service_date_str = request.form.get('service_date', date.today().strftime('%Y-%m-%d'))
    service_type = request.form.get('service_type', 'Sunday Service')

    if not full_name:
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return jsonify({'success': False, 'message': 'Full name is required'}), 400
        flash('Full name is required to register attendee.', 'danger')
        return redirect(url_for('attendance.checkin', date=service_date_str, service_type=service_type))

    last_member = Member.query.order_by(Member.id.desc()).first()
    next_num = (last_member.id + 1) if last_member else 1
    member_code = f"NHGCC-{next_num:04d}"

    parts = full_name.split(maxsplit=1)
    first_name = parts[0]
    last_name = parts[1] if len(parts) > 1 else ''

    new_member = Member(
        member_code=member_code,
        full_name=full_name,
        first_name=first_name,
        last_name=last_name,
        phone=phone,
        email=email,
        gender=gender,
        address=address,
        status=status,
        department=department,
        date_joined=date.today(),
        is_active=True
    )
    db.session.add(new_member)
    db.session.flush()

    try:
        service_date = datetime.strptime(service_date_str, '%Y-%m-%d').date()
    except ValueError:
        service_date = date.today()

    att = Attendance(
        member_id=new_member.id,
        service_date=service_date,
        service_type=service_type,
        status='Present',
        check_in_time=datetime.now(),
        marked_by='Quick Entrance Registration'
    )
    db.session.add(att)
    db.session.commit()

    if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        return jsonify({
            'success': True,
            'message': f"Successfully registered and checked in {new_member.full_name}!",
            'member': new_member.to_dict()
        })

    flash(f"New member '{new_member.full_name}' successfully added and marked present for today's service!", 'success')
    return redirect(url_for('attendance.checkin', date=service_date_str, service_type=service_type))


@attendance_bp.route('/history')
def history():
    """Displays chronological history of past Sunday attendances."""
    import traceback
    try:
        records = (
            db.session.query(
                Attendance.service_date,
                Attendance.service_type,
                func.count(Attendance.id).label('total_present')
            )
            .filter(Attendance.status == 'Present')
            .group_by(Attendance.service_date, Attendance.service_type)
            .order_by(Attendance.service_date.desc())
            .all()
        )
    except Exception as e:
        print(f"[NHGCC Error] Attendance history query error: {e}\n{traceback.format_exc()}")
        records = []

    formatted_records = []
    for rec in records:
        try:
            s_date = rec[0] if isinstance(rec, (tuple, list)) else getattr(rec, 'service_date', None)
            s_type = rec[1] if isinstance(rec, (tuple, list)) else getattr(rec, 'service_type', 'Sunday Service')
            count = rec[2] if isinstance(rec, (tuple, list)) else getattr(rec, 'total_present', 0)

            if isinstance(s_date, str):
                try:
                    s_date = datetime.strptime(s_date, '%Y-%m-%d').date()
                except Exception:
                    pass

            date_formatted = s_date.strftime('%d %B %Y') if hasattr(s_date, 'strftime') else str(s_date)
            date_iso = s_date.strftime('%Y-%m-%d') if hasattr(s_date, 'strftime') else str(s_date)
            day_name = s_date.strftime('%A') if hasattr(s_date, 'strftime') else 'Sunday'

            formatted_records.append({
                'service_date': s_date,
                'service_date_formatted': date_formatted,
                'service_date_iso': date_iso,
                'day_name': day_name,
                'service_type': s_type or 'Sunday Service',
                'total_present': int(count) if count else 0
            })
        except Exception as row_err:
            print(f"[NHGCC Warning] Error formatting row: {row_err}")

    return render_template('attendance/history.html', history_records=formatted_records)


@attendance_bp.route('/export')
def export_excel():
    """Exports a clean Excel register of Present attendees, Absent members, or both."""
    date_str = request.args.get('date')
    service_type = request.args.get('service_type', 'Sunday Service')
    scope = request.args.get('scope', 'both').lower()  # 'both', 'present', or 'absent'

    try:
        service_date = datetime.strptime(date_str, '%Y-%m-%d').date()
    except Exception:
        service_date = get_latest_sunday()

    # 1. Present attendees
    present_attendances = (
        Attendance.query.filter_by(
            service_date=service_date,
            service_type=service_type,
            status='Present'
        )
        .join(Member)
        .order_by(Member.full_name)
        .all()
    )
    present_ids = {a.member_id: a for a in present_attendances}

    # 2. All active members & absent list
    all_members = Member.query.filter_by(is_active=True).order_by(Member.full_name).all()
    absent_members = [m for m in all_members if m.id not in present_ids]

    # Query last attended date for absent members
    prior_attendances = (
        Attendance.query.filter(
            Attendance.service_date < service_date,
            Attendance.status == 'Present'
        )
        .order_by(Attendance.service_date.desc())
        .all()
    )
    last_attended_map = {}
    for a in prior_attendances:
        if a.member_id not in last_attended_map:
            last_attended_map[a.member_id] = a.service_date

    thin_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    wb = openpyxl.Workbook()

    # ----------------- TAB: PRESENT ATTENDEES -----------------
    def build_present_sheet(ws):
        ws.title = f"Present ({len(present_attendances)})"
        ws.views.sheetView[0].showGridLines = True
        ws.append(["NHGCC OSHODI • ATTENDANCE REGISTER (PRESENT)"])
        ws.append([f"Service: {service_type} | Date: {service_date.strftime('%A, %d %B %Y')} | Total Present: {len(present_attendances)}"])
        ws.append([])

        ws['A1'].font = Font(name='Segoe UI', size=15, bold=True, color='065F46')
        ws['A2'].font = Font(name='Segoe UI', size=10.5, italic=True)

        headers = ['S/N', 'Member Code', 'Full Name', 'Gender', 'Phone Number', 'Email', 'Department', 'Member Status', 'Check-In Time']
        ws.append(headers)

        header_fill = PatternFill(start_color='065F46', end_color='065F46', fill_type='solid')
        header_font = Font(name='Segoe UI', size=10.5, bold=True, color='FFFFFF')

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=4, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center' if col_idx in (1, 2, 4, 8, 9) else 'left')
            cell.border = cell_border

        for idx, a in enumerate(present_attendances, 1):
            m = a.member
            check_time_str = a.check_in_time.strftime('%I:%M %p') if a.check_in_time else 'Present'
            row_vals = [
                idx,
                m.member_code or f"NHGCC-{m.id:04d}",
                m.display_title_name,
                m.gender or 'Not Specified',
                m.phone or '—',
                m.email or '—',
                m.department or 'Congregation',
                m.status or 'Active',
                check_time_str
            ]
            ws.append(row_vals)
            row_idx = idx + 4
            for col_idx in range(1, len(row_vals) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.border = cell_border
                c.font = Font(name='Segoe UI', size=10)
                if idx % 2 == 0:
                    c.fill = PatternFill(start_color='F0FDF4', end_color='F0FDF4', fill_type='solid')
                if col_idx in (1, 2, 4, 8, 9):
                    c.alignment = Alignment(horizontal='center')

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # ----------------- TAB: ABSENT MEMBERS -----------------
    def build_absent_sheet(ws):
        ws.title = f"Absent ({len(absent_members)})"
        ws.views.sheetView[0].showGridLines = True
        ws.append(["NHGCC OSHODI • LIST OF MEMBERS WHO WERE ABSENT"])
        ws.append([f"Service: {service_type} | Date: {service_date.strftime('%A, %d %B %Y')} | Total Absent: {len(absent_members)}"])
        ws.append([])

        ws['A1'].font = Font(name='Segoe UI', size=15, bold=True, color='991B1B')
        ws['A2'].font = Font(name='Segoe UI', size=10.5, italic=True)

        headers = ['S/N', 'Member Code', 'Full Name', 'Gender', 'Phone Number', 'Email', 'Department', 'Member Status', 'Last Attended Date', 'Follow-up Status']
        ws.append(headers)

        header_fill = PatternFill(start_color='991B1B', end_color='991B1B', fill_type='solid')
        header_font = Font(name='Segoe UI', size=10.5, bold=True, color='FFFFFF')

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=4, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center' if col_idx in (1, 2, 4, 8, 9, 10) else 'left')
            cell.border = cell_border

        for idx, m in enumerate(absent_members, 1):
            last_dt = last_attended_map.get(m.id)
            last_dt_str = last_dt.strftime('%d-%b-%Y') if last_dt else 'No Prior Attendance'
            row_vals = [
                idx,
                m.member_code or f"NHGCC-{m.id:04d}",
                m.display_title_name,
                m.gender or 'Not Specified',
                m.phone or '—',
                m.email or '—',
                m.department or 'Congregation',
                m.status or 'Active',
                last_dt_str,
                'Needs Follow-Up'
            ]
            ws.append(row_vals)
            row_idx = idx + 4
            for col_idx in range(1, len(row_vals) + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.border = cell_border
                c.font = Font(name='Segoe UI', size=10)
                if idx % 2 == 0:
                    c.fill = PatternFill(start_color='FEF2F2', end_color='FEF2F2', fill_type='solid')
                if col_idx in (1, 2, 4, 8, 9, 10):
                    c.alignment = Alignment(horizontal='center')

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # ----------------- TAB: SUMMARY -----------------
    def build_summary_sheet(ws):
        ws.title = "Executive Summary"
        ws.views.sheetView[0].showGridLines = True
        ws.append(["NATIONAL HOLY GHOST CHURCH OF CHRIST (NHGCC) OSHODI"])
        ws.append([f"SUNDAY ATTENDANCE & ABSENTEE REGISTER • {service_date.strftime('%d %B %Y').upper()}"])
        ws.append([])

        ws['A1'].font = Font(name='Segoe UI', size=16, bold=True, color='0B192C')
        ws['A2'].font = Font(name='Segoe UI', size=11, bold=True, color='D97706')

        ws.append(["METRIC", "VALUE", "NOTES"])
        for col in range(1, 4):
            c = ws.cell(row=4, column=col)
            c.fill = PatternFill(start_color='0B192C', end_color='0B192C', fill_type='solid')
            c.font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
            c.alignment = Alignment(horizontal='center' if col == 2 else 'left')

        summary_rows = [
            ("Service Date", service_date.strftime('%A, %d %B %Y'), "Official Sunday Service"),
            ("Service Type", service_type, "Entrance Desk Recording"),
            ("Total Registered Parish Members", len(all_members), "Active database membership"),
            ("Members Present (Who Came)", len(present_attendances), f"{len(present_attendances)/len(all_members)*100:.1f}% Attendance Rate" if all_members else "0%"),
            ("Members Absent (Who Missed)", len(absent_members), f"{len(absent_members)/len(all_members)*100:.1f}% Absence Rate" if all_members else "0%"),
            ("Male Attendees Present", sum(1 for a in present_attendances if a.member.gender == 'Male'), ""),
            ("Female Attendees Present", sum(1 for a in present_attendances if a.member.gender == 'Female'), ""),
            ("Unspecified Gender", sum(1 for a in present_attendances if a.member.gender not in ('Male', 'Female')), "")
        ]

        for idx, (label, val, note) in enumerate(summary_rows, 5):
            ws.append([label, val, note])
            for col in range(1, 4):
                cell = ws.cell(row=idx, column=col)
                cell.border = cell_border
                cell.font = Font(name='Segoe UI', size=10, bold=(col == 2))
                if col == 2:
                    cell.alignment = Alignment(horizontal='center')
                    if label == "Members Present (Who Came)":
                        cell.fill = PatternFill(start_color='D1FAE5', end_color='D1FAE5', fill_type='solid')
                        cell.font = Font(name='Segoe UI', size=11, bold=True, color='065F46')
                    elif label == "Members Absent (Who Missed)":
                        cell.fill = PatternFill(start_color='FEE2E2', end_color='FEE2E2', fill_type='solid')
                        cell.font = Font(name='Segoe UI', size=11, bold=True, color='991B1B')

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    date_stamp = service_date.strftime('%Y%m%d')

    if scope == 'present':
        ws = wb.active
        build_present_sheet(ws)
        filename = f"NHGCC_Present_Attendees_{date_stamp}.xlsx"
    elif scope == 'absent':
        ws = wb.active
        build_absent_sheet(ws)
        filename = f"NHGCC_Absent_Members_{date_stamp}.xlsx"
    else:
        # Default: Full combined workbook with Summary, Present, and Absent tabs
        ws_sum = wb.active
        build_summary_sheet(ws_sum)
        ws_pres = wb.create_sheet()
        build_present_sheet(ws_pres)
        ws_abs = wb.create_sheet()
        build_absent_sheet(ws_abs)
        filename = f"NHGCC_Attendance_Register_Full_{date_stamp}.xlsx"

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
